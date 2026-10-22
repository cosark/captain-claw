"""Exporting a Data library table as CSV, JSON or an Excel workbook.

The kiosk has no Datastore panel any more, so the Data library's export is
the one way out for a table: ``format=csv`` (the default, as before),
``json`` (one object per row) or ``xlsx``. Every format writes the table's
visible columns only — never the datastore's hidden creator columns — a
text cell that looks like a formula stays text in the workbook, anything
else is a 400, and another user's table is refused whatever the format.

In the CSV, a cell a member of a shared agent stored that starts like a
formula gets a leading ``'`` (J13), as the agent's own export does; the
owner's cells, and every cell of a table with no creator column, stay as
stored.
"""

from __future__ import annotations

import csv
import io
import json
import sqlite3
from contextlib import closing
from pathlib import Path

import httpx
import openpyxl
import pytest

from captain_claw.datastore import DatastoreManager
from captain_claw.flight_deck import auth as fd_auth
from captain_claw.flight_deck import data_library as dl
from captain_claw.flight_deck import rate_limiter, server
from captain_claw.flight_deck.auth import create_access_token, set_auth_db
from captain_claw.flight_deck.db import FlightDeckDB

BOB = "user-bob"
CAROL = "user-carol"
DANA = "user-dana"  # a member of Bob's shared agent
FORMULA = '=HYPERLINK("http://evil.example","click")'
ROWS = [
    {"name": "Ana", "email": "ana@x.co", "score": 7, "note": FORMULA},
    {"name": "Ivo", "email": None, "score": 3, "note": "Žuta kuća"},
]


@pytest.fixture
async def deck(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """Auth-enabled deck: Bob's stopped agent keeps a ``contacts`` table;
    Carol owns an agent of her own. Yields the table's library id."""
    prev = fd_auth._db
    db = FlightDeckDB(tmp_path / "fd.db")
    await db.init()
    set_auth_db(db)
    for uid in (BOB, CAROL):
        await db._db.execute(
            "INSERT INTO users (id, email, password_hash, display_name, role, created_at, updated_at)"
            " VALUES (?, ?, 'h', ?, 'user', '2026-01-01T00:00:00Z', '2026-01-01T00:00:00Z')",
            (uid, f"{uid}@x.co", uid))
    await db._db.commit()

    data = tmp_path / "fd-data"
    data.mkdir()
    monkeypatch.setattr(server, "DATA_DIR", data)
    monkeypatch.setattr(server, "PROCESS_REGISTRY_FILE", data / ".processes.json")
    monkeypatch.setattr(server, "_processes", {})
    monkeypatch.setattr(server, "AUTH_ENABLED", True)
    monkeypatch.setenv("FD_AUTH_ENABLED", "true")
    monkeypatch.setenv("FD_ALLOWED_HOSTS", "fd.test")
    monkeypatch.delenv("FD_LOCKDOWN", raising=False)
    monkeypatch.setattr(rate_limiter, "_api_limiter", rate_limiter._SlidingWindow())
    monkeypatch.setattr(dl, "_STORE", dl.DataStore(tmp_path / "data_library.db"))

    server._save_process_registry({
        "bob-agent": {"slug": "bob-agent", "name": "Bob's agent", "web_port": 24101,
                      "web_auth": "bob-tok", "owner": BOB, "pid": None},
        "carol-agent": {"slug": "carol-agent", "name": "Carol's agent", "web_port": 24102,
                        "web_auth": "carol-tok", "owner": CAROL, "pid": None},
    })
    mgr = DatastoreManager(
        db_path=data / "bob-agent" / "data" / "home-config-parent" / ".captain-claw" / "datastore.db")
    await mgr.create_table("contacts", [
        {"name": "name", "type": "text"}, {"name": "email", "type": "text"},
        {"name": "score", "type": "integer"}, {"name": "note", "type": "text"},
    ])
    await mgr.insert_rows("contacts", ROWS)
    await mgr.close()
    # What a scan does before it asks the model for summaries.
    dl.get_store().sync_agent("bob-agent", dl.scan_agent("bob-agent", {}))
    try:
        yield dl.DataStore.table_id("bob-agent", "contacts")
    finally:
        await db.close()
        fd_auth._db = prev


def _client() -> httpx.AsyncClient:
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=server.app, client=("203.0.113.9", 40001)), base_url="http://fd.test")


def _auth(uid: str) -> dict:
    return {"Authorization": f"Bearer {create_access_token(uid)}"}


async def _export(table_id: str, uid: str = BOB, fmt: str | None = None) -> httpx.Response:
    params = {"format": fmt} if fmt else {}
    async with _client() as c:
        return await c.get(f"/fd/data-library/tables/{table_id}/export", params=params, headers=_auth(uid))


async def test_csv_stays_the_default(deck: str):
    r = await _export(deck)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    assert 'filename="contacts.csv"' in r.headers["content-disposition"]
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "name,email,score,note"
    assert lines[2] == "Ivo,,3,Žuta kuća"


async def test_json_is_one_object_per_row(deck: str):
    r = await _export(deck, fmt="json")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/json")
    assert 'filename="contacts.json"' in r.headers["content-disposition"]
    # The visible columns only, values as stored (a missing one is null).
    assert json.loads(r.content) == ROWS


async def test_xlsx_is_a_workbook_with_the_rows(deck: str):
    r = await _export(deck, fmt="xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="contacts.xlsx"' in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["contacts"]
    ws = wb["contacts"]
    assert [list(row) for row in ws.iter_rows(values_only=True)] == [
        ["name", "email", "score", "note"],
        ["Ana", "ana@x.co", 7, FORMULA],
        ["Ivo", None, 3, "Žuta kuća"],
    ]
    # A cell that looks like a formula is written as text, so Excel never runs it.
    assert ws["D2"].data_type == "s"
    assert ws["A1"].font.bold


async def test_an_unknown_format_is_refused(deck: str):
    r = await _export(deck, fmt="pdf")
    assert r.status_code == 400
    assert "csv, json or xlsx" in r.json()["detail"]


@pytest.mark.parametrize("fmt", [None, "csv", "json", "xlsx"])
async def test_another_users_table_is_refused(deck: str, fmt: str | None):
    r = await _export(deck, uid=CAROL, fmt=fmt)
    assert r.status_code == 404
    assert "Ana" not in r.text


def _write(*statements: tuple) -> None:
    """Change Bob's agent's datastore file straight on disk, as the agent would."""
    path = dl.datastore_path("bob-agent")
    assert path is not None
    with closing(sqlite3.connect(path)) as c:
        for sql, *args in statements:
            c.execute(sql, *args)
        c.commit()


def _add_member_row() -> None:
    """A row Dana stored through Bob's shared agent: it carries her creator stamp."""
    _write(('INSERT INTO ds_contacts (name, email, score, note, "_created_by", "_created_by_name")'
            " VALUES (?, ?, ?, ?, ?, ?)", ("Dana", "@SUM(A1:A9)", -4, FORMULA, DANA, "Dana")))


def _csv_rows(r: httpx.Response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))


async def test_csv_defuses_a_members_formula_but_not_the_owners(deck: str):
    _add_member_row()
    r = await _export(deck, fmt="csv")
    assert r.status_code == 200, r.text
    rows = _csv_rows(r)
    # The creator columns are read, never written.
    assert rows[0] == ["name", "email", "score", "note"]
    # Bob's own cell is exported as he stored it.
    assert rows[1] == ["Ana", "ana@x.co", "7", FORMULA]
    assert rows[2] == ["Ivo", "", "3", "Žuta kuća"]
    # Dana's text cells that start like a formula get a leading "'"; her
    # number and her plain text stay as they are.
    assert rows[3] == ["Dana", "'@SUM(A1:A9)", "-4", "'" + FORMULA]


async def test_json_and_xlsx_keep_a_members_cell_as_stored(deck: str):
    _add_member_row()
    r = await _export(deck, fmt="json")
    assert r.status_code == 200, r.text
    # Nothing runs JSON, so it carries the cell exactly.
    assert json.loads(r.content)[2] == {"name": "Dana", "email": "@SUM(A1:A9)", "score": -4, "note": FORMULA}

    r = await _export(deck, fmt="xlsx")
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["contacts"]
    assert [c.value for c in ws[4]] == ["Dana", "@SUM(A1:A9)", -4, FORMULA]
    # The workbook writes it as text instead, so Excel never runs it.
    assert ws["B4"].data_type == "s"
    assert ws["D4"].data_type == "s"


async def test_a_table_with_no_creator_column_exports_as_stored(deck: str):
    """A table from before shared agents has no creator stamp: every row is
    the owner's, as the agent's own export reads it."""
    path = dl.datastore_path("bob-agent")
    assert path is not None
    mgr = DatastoreManager(db_path=path)
    await mgr.create_table("legacy", [{"name": "k", "type": "text"}, {"name": "n", "type": "integer"}])
    await mgr.insert_rows("legacy", [{"k": FORMULA, "n": 1}, {"k": "-2+3", "n": -2}])
    await mgr.close()
    _write(('DROP INDEX "ix_ds_legacy_created_by"',),
           ('ALTER TABLE ds_legacy DROP COLUMN "_created_by"',),
           ('ALTER TABLE ds_legacy DROP COLUMN "_created_by_name"',))
    dl.get_store().sync_agent("bob-agent", dl.scan_agent("bob-agent", {}))

    r = await _export(dl.DataStore.table_id("bob-agent", "legacy"), fmt="csv")
    assert r.status_code == 200, r.text
    assert _csv_rows(r) == [["k", "n"], [FORMULA, "1"], ["-2+3", "-2"]]
